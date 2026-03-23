import base64
import io
import logging
import uuid
from datetime import datetime, timezone
from typing import List, Optional
from uuid import UUID

from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile, status
from pydantic import BaseModel

from app.api.deps import get_current_user, get_tenant_id
from app.db.models import User, UserRole
from mcp_server.tools.chroma import (
    chroma_delete_coach_doc,
    chroma_delete_methodology_doc,
    chroma_ingest_coach_docs,
    chroma_ingest_methodology_docs,
    chroma_list_coach_docs,
    chroma_list_methodology_docs,
    chroma_query_coach_docs,
)
from mcp_server.tools.postgres import (
    pg_archive_coach_document,
    pg_archive_methodology_document,
    pg_get_coach_document,
    pg_get_methodology_document,
    pg_get_user_by_id,
    pg_list_active_coach_documents,
    pg_get_users_by_ids,
    pg_save_coach_document,
    pg_save_methodology_document,
    pg_update_methodology_document,
)

router = APIRouter(prefix="/embeddings", tags=["embeddings"])
logger = logging.getLogger(__name__)


# ─────────────────────────────── Schemas ──────────────────────────────────────

class IngestRequest(BaseModel):
    documents: List[str]
    ids: Optional[List[str]] = None
    metadatas: Optional[List[dict]] = None


class IngestResponse(BaseModel):
    ingested: int
    collection: str


class DocumentInfo(BaseModel):
    doc_id: str
    filename: str
    chunk_count: int
    uploaded_at: Optional[str] = None


class CoachDocumentInfo(BaseModel):
    doc_id: str
    filename: str
    client_id: str
    client_email: Optional[str] = None
    chunk_count: int
    uploaded_at: Optional[str] = None


# ─────────────────────────────── Auth helpers ─────────────────────────────────

def _require_admin(current_user: User) -> None:
    if current_user.role != UserRole.admin:
        logger.warning(
            "Rejected admin-only document operation user_id=%s role=%s",
            current_user.id,
            current_user.role.value,
        )
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Only admins can manage methodology documents",
        )


def _require_coach(current_user: User) -> None:
    if current_user.role != UserRole.coach:
        logger.warning(
            "Rejected coach-only document operation user_id=%s role=%s",
            current_user.id,
            current_user.role.value,
        )
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Only coaches can manage client conversation documents",
        )


# ─────────────────────────────── File helpers ─────────────────────────────────

def _extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from PDF, DOCX, or TXT/MD files."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            logger.exception("PDF parsing failed filename=%s", filename)
            raise HTTPException(status_code=400, detail=f"Failed to parse PDF: {exc}")
    elif lower.endswith(".docx"):
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception as exc:
            logger.exception("DOCX parsing failed filename=%s", filename)
            raise HTTPException(status_code=400, detail=f"Failed to parse DOCX: {exc}")
    elif lower.endswith(".xlsx"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    line = "\t".join("" if v is None else str(v) for v in row)
                    if line.strip():
                        lines.append(line)
            return "\n".join(lines)
        except Exception as exc:
            logger.exception("XLSX parsing failed filename=%s", filename)
            raise HTTPException(status_code=400, detail=f"Failed to parse XLSX: {exc}")
    elif lower.endswith(".txt") or lower.endswith(".md"):
        return file_bytes.decode("utf-8", errors="replace")
    else:
        logger.warning("Rejected unsupported document type filename=%s", filename)
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Accepted: PDF, DOCX, XLSX, TXT, MD",
        )


def _split_chunks(text: str, chunk_size: int = 1000, overlap: int = 100) -> List[str]:
    chunks, start = [], 0
    while start < len(text):
        chunks.append(text[start : start + chunk_size].strip())
        start += chunk_size - overlap
    return [c for c in chunks if c]


# ─────────────────────────────── Coach helpers ────────────────────────────────

async def _verify_coach_client(coach: User, client_id: UUID) -> dict:
    """Verify client belongs to this coach. Returns the client dict."""
    client = await pg_get_user_by_id(str(client_id))
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    if client["tenant_id"] != str(coach.tenant_id):
        raise HTTPException(status_code=403, detail="Client is in a different tenant")
    if client["role"] != "client":
        raise HTTPException(status_code=400, detail="The specified user is not a client")
    if client.get("coach_id") != str(coach.id):
        raise HTTPException(status_code=403, detail="You are not assigned to this client")
    return client


async def _get_active_coach_doc(coach: User, doc_id: str) -> dict:
    """Load an active CoachDocument owned by this coach or raise 404/403."""
    doc = await pg_get_coach_document(doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    if doc["coach_id"] != str(coach.id):
        raise HTTPException(status_code=403, detail="You do not own this document")
    return doc


# ═══════════════════════════════════════════════════════════════════════════════
# ADMIN — methodology documents
# ═══════════════════════════════════════════════════════════════════════════════

@router.post("/ingest", response_model=IngestResponse)
async def ingest_documents_endpoint(
    payload: IngestRequest,
    current_user: User = Depends(get_current_user),
    tenant_id: str = Depends(get_tenant_id),
) -> IngestResponse:
    """Ingest raw text documents into the tenant's ChromaDB vector store. Admin only."""
    _require_admin(current_user)

    if not payload.documents:
        raise HTTPException(status_code=400, detail="No documents provided")

    doc_ids = payload.ids or [str(uuid.uuid4()) for _ in payload.documents]
    if len(doc_ids) != len(payload.documents):
        raise HTTPException(status_code=400, detail="Number of IDs must match number of documents")

    try:
        chroma_ingest_methodology_docs(
            tenant_id=tenant_id,
            documents=payload.documents,
            ids=doc_ids,
            metadatas=payload.metadatas,
        )
    except Exception:
        logger.exception(
            "Raw document ingestion failed tenant_id=%s user_id=%s document_count=%s",
            tenant_id,
            current_user.id,
            len(payload.documents),
        )
        raise HTTPException(status_code=500, detail="Document ingestion failed. Check backend logs.")

    return IngestResponse(
        ingested=len(payload.documents), collection=f"tenant_{tenant_id}_resources"
    )


@router.post("/upload", response_model=IngestResponse)
async def upload_document(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    tenant_id: str = Depends(get_tenant_id),
) -> IngestResponse:
    """Upload a methodology document (PDF, DOCX, TXT) into the RAG store. Admin only."""
    _require_admin(current_user)

    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename provided")

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")

    text = _extract_text(file_bytes, file.filename)
    if not text.strip():
        raise HTTPException(status_code=400, detail="No text could be extracted from the file")

    chunks = _split_chunks(text)
    doc_id = str(uuid.uuid4())
    uploaded_at = datetime.now(timezone.utc).isoformat()
    chunk_ids = [f"{doc_id}_{i}" for i in range(len(chunks))]
    metadatas = [
        {
            "doc_id": doc_id,
            "filename": file.filename,
            "uploaded_by": str(current_user.id),
            "uploaded_at": uploaded_at,
            "chunk_index": i,
        }
        for i in range(len(chunks))
    ]

    try:
        chroma_ingest_methodology_docs(
            tenant_id=tenant_id, documents=chunks, ids=chunk_ids, metadatas=metadatas
        )
    except Exception:
        logger.exception(
            "Methodology upload ingestion failed tenant_id=%s user_id=%s filename=%s",
            tenant_id,
            current_user.id,
            file.filename,
        )
        raise HTTPException(status_code=500, detail="Document upload failed during indexing.")

    logger.info(
        "Admin %s uploaded methodology doc '%s' (%d chunks) tenant=%s",
        current_user.id,
        file.filename,
        len(chunks),
        tenant_id,
    )
    await pg_save_methodology_document(
        doc_id=doc_id,
        tenant_id=tenant_id,
        filename=file.filename,
        content_type=file.content_type or "application/octet-stream",
        file_bytes_b64=base64.b64encode(file_bytes).decode(),
        uploaded_by=str(current_user.id),
        uploaded_at=uploaded_at,
    )
    return IngestResponse(
        ingested=len(chunks), collection=f"tenant_{tenant_id}_resources"
    )


@router.get("/documents", response_model=List[DocumentInfo])
async def list_uploaded_documents(
    current_user: User = Depends(get_current_user),
    tenant_id: str = Depends(get_tenant_id),
) -> List[DocumentInfo]:
    """List all methodology documents for this tenant. Admin only."""
    _require_admin(current_user)
    try:
        docs = chroma_list_methodology_docs(tenant_id=tenant_id)
    except Exception:
        logger.exception(
            "Document listing failed tenant_id=%s user_id=%s", tenant_id, current_user.id
        )
        raise HTTPException(status_code=500, detail="Document listing failed.")
    logger.info(
        "Listed methodology docs tenant_id=%s user_id=%s count=%s",
        tenant_id,
        current_user.id,
        len(docs),
    )
    return [DocumentInfo(**d) for d in docs]


@router.get("/documents/{doc_id}/download")
async def download_methodology_document(
    doc_id: str,
    current_user: User = Depends(get_current_user),
    tenant_id: str = Depends(get_tenant_id),
) -> Response:
    """Download the original file for a methodology document. Admin only."""
    _require_admin(current_user)
    doc = await pg_get_methodology_document(doc_id=doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found or has been archived")
    file_bytes = base64.b64decode(doc["file_bytes_b64"])
    return Response(
        content=file_bytes,
        media_type=doc["content_type"],
        headers={"Content-Disposition": f'attachment; filename="{doc["filename"]}"'},
    )


@router.put("/documents/{doc_id}", response_model=IngestResponse)
async def replace_methodology_document(
    doc_id: str,
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    tenant_id: str = Depends(get_tenant_id),
) -> IngestResponse:
    """Replace an existing methodology document with a new file. Admin only."""
    _require_admin(current_user)
    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename provided")

    existing = await pg_get_methodology_document(doc_id=doc_id)
    if not existing:
        raise HTTPException(status_code=404, detail="Document not found or has been archived")

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")

    text = _extract_text(file_bytes, file.filename)
    if not text.strip():
        raise HTTPException(status_code=400, detail="No text could be extracted from the file")

    try:
        chroma_delete_methodology_doc(tenant_id=tenant_id, doc_id=doc_id)
    except Exception:
        logger.exception(
            "Replace: chroma delete failed tenant_id=%s doc_id=%s", tenant_id, doc_id
        )
        raise HTTPException(status_code=500, detail="Failed to remove old document from index.")

    uploaded_at = datetime.now(timezone.utc).isoformat()
    chunks = _split_chunks(text)
    chunk_ids = [f"{doc_id}_{i}" for i in range(len(chunks))]
    metadatas = [
        {
            "doc_id": doc_id,
            "filename": file.filename,
            "uploaded_by": str(current_user.id),
            "uploaded_at": uploaded_at,
            "chunk_index": i,
        }
        for i in range(len(chunks))
    ]
    try:
        chroma_ingest_methodology_docs(
            tenant_id=tenant_id, documents=chunks, ids=chunk_ids, metadatas=metadatas
        )
    except Exception:
        logger.exception(
            "Replace: chroma ingest failed tenant_id=%s doc_id=%s filename=%s",
            tenant_id,
            doc_id,
            file.filename,
        )
        raise HTTPException(status_code=500, detail="Failed to index replacement document.")

    await pg_update_methodology_document(
        doc_id=doc_id,
        filename=file.filename,
        content_type=file.content_type or "application/octet-stream",
        file_bytes_b64=base64.b64encode(file_bytes).decode(),
        uploaded_by=str(current_user.id),
        uploaded_at=uploaded_at,
    )
    logger.info(
        "Admin %s replaced methodology doc doc_id=%s new_filename='%s' tenant=%s",
        current_user.id,
        doc_id,
        file.filename,
        tenant_id,
    )
    return IngestResponse(ingested=len(chunks), collection=f"tenant_{tenant_id}_resources")


@router.delete("/documents/{doc_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_uploaded_document(
    doc_id: str,
    current_user: User = Depends(get_current_user),
    tenant_id: str = Depends(get_tenant_id),
) -> None:
    """Delete a methodology document from the RAG store. Admin only."""
    _require_admin(current_user)
    try:
        deleted = chroma_delete_methodology_doc(tenant_id=tenant_id, doc_id=doc_id)
    except Exception:
        logger.exception(
            "Document delete failed tenant_id=%s doc_id=%s", tenant_id, doc_id
        )
        raise HTTPException(status_code=500, detail="Document delete failed.")
    if not deleted:
        raise HTTPException(status_code=404, detail="Document not found")
    await pg_archive_methodology_document(doc_id=doc_id)
    logger.info(
        "Deleted methodology doc tenant_id=%s user_id=%s doc_id=%s",
        tenant_id,
        current_user.id,
        doc_id,
    )


# ═══════════════════════════════════════════════════════════════════════════════
# COACH — client conversation documents
# ═══════════════════════════════════════════════════════════════════════════════

@router.post("/coach/upload", response_model=IngestResponse)
async def upload_coach_document(
    file: UploadFile = File(...),
    client_id: UUID = Form(...),
    current_user: User = Depends(get_current_user),
    tenant_id: str = Depends(get_tenant_id),
) -> IngestResponse:
    """Coach uploads a past-conversation document for an assigned client."""
    _require_coach(current_user)
    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename provided")

    client = await _verify_coach_client(current_user, client_id)

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")

    content_type = file.content_type or "application/octet-stream"
    text = _extract_text(file_bytes, file.filename)
    if not text.strip():
        raise HTTPException(status_code=400, detail="No text could be extracted from the file")

    chunks = _split_chunks(text)
    doc_id = str(uuid.uuid4())
    uploaded_at = datetime.now(timezone.utc).isoformat()
    chunk_ids = [f"{doc_id}_{i}" for i in range(len(chunks))]
    metadatas = [
        {
            "doc_id": doc_id,
            "filename": file.filename,
            "coach_id": str(current_user.id),
            "client_id": str(client["id"]),
            "uploaded_at": uploaded_at,
            "chunk_index": i,
        }
        for i in range(len(chunks))
    ]

    try:
        chroma_ingest_coach_docs(
            tenant_id=tenant_id, documents=chunks, ids=chunk_ids, metadatas=metadatas
        )
    except Exception:
        logger.exception(
            "Coach document ingestion failed tenant_id=%s coach_id=%s client_id=%s filename=%s",
            tenant_id,
            current_user.id,
            client_id,
            file.filename,
        )
        raise HTTPException(status_code=500, detail="Document upload failed during indexing.")

    # Persist original bytes via MCP tool
    await pg_save_coach_document(
        doc_id=doc_id,
        coach_id=str(current_user.id),
        client_id=str(client["id"]),
        tenant_id=str(current_user.tenant_id),
        filename=file.filename,
        content_type=content_type,
        file_bytes_b64=base64.b64encode(file_bytes).decode(),
        uploaded_at=uploaded_at,
    )

    logger.info(
        "Coach %s uploaded client doc '%s' (%d chunks) for client %s tenant=%s",
        current_user.id,
        file.filename,
        len(chunks),
        client_id,
        tenant_id,
    )
    return IngestResponse(
        ingested=len(chunks), collection=f"tenant_{tenant_id}_coach_conversations"
    )


@router.get("/coach/documents", response_model=List[CoachDocumentInfo])
async def list_coach_uploaded_documents(
    client_id: Optional[UUID] = None,
    current_user: User = Depends(get_current_user),
    tenant_id: str = Depends(get_tenant_id),
) -> List[CoachDocumentInfo]:
    """List active coach-conversation documents (optionally filtered by client). Coach only."""
    _require_coach(current_user)

    client_filter_str = str(client_id) if client_id else None
    chroma_docs = chroma_list_coach_docs(
        tenant_id=tenant_id,
        coach_id=str(current_user.id),
        client_id=client_filter_str,
    )

    # Resolve unique client emails in one batch call via MCP tool
    client_ids_needed = list({d["client_id"] for d in chroma_docs if d.get("client_id")})
    client_emails: dict = {}
    if client_ids_needed:
        try:
            users = await pg_get_users_by_ids(client_ids_needed)
            client_emails = {u["id"]: u["email"] for u in users}
        except Exception:
            logger.warning("Could not resolve client emails for coach document list")

    result = [
        CoachDocumentInfo(
            doc_id=d["doc_id"],
            filename=d["filename"],
            client_id=d["client_id"],
            client_email=client_emails.get(d["client_id"]),
            chunk_count=d["chunk_count"],
            uploaded_at=d.get("uploaded_at"),
        )
        for d in chroma_docs
    ]
    logger.info(
        "Listed coach docs coach_id=%s tenant_id=%s count=%s",
        current_user.id,
        tenant_id,
        len(result),
    )
    return result


@router.get("/coach/documents/{doc_id}/download")
async def download_coach_document(
    doc_id: str,
    current_user: User = Depends(get_current_user),
    tenant_id: str = Depends(get_tenant_id),
) -> Response:
    """Return the original uploaded file bytes for download. Coach only (own docs)."""
    _require_coach(current_user)
    doc = await _get_active_coach_doc(current_user, doc_id)
    file_bytes = base64.b64decode(doc["file_bytes_b64"])
    logger.info(
        "Coach %s downloaded doc_id=%s filename=%s", current_user.id, doc_id, doc["filename"]
    )
    return Response(
        content=file_bytes,
        media_type=doc["content_type"],
        headers={"Content-Disposition": f'attachment; filename="{doc["filename"]}"'},
    )


@router.put("/coach/documents/{doc_id}", response_model=IngestResponse)
async def replace_coach_document(
    doc_id: str,
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    tenant_id: str = Depends(get_tenant_id),
) -> IngestResponse:
    """Hard-replace a coach document: archive old record, delete old Chroma chunks, ingest new. Coach only."""
    _require_coach(current_user)
    old_doc = await _get_active_coach_doc(current_user, doc_id)

    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename provided")

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")

    content_type = file.content_type or "application/octet-stream"
    text = _extract_text(file_bytes, file.filename)
    if not text.strip():
        raise HTTPException(status_code=400, detail="No text could be extracted from the file")

    chunks = _split_chunks(text)
    new_doc_id = str(uuid.uuid4())
    uploaded_at = datetime.now(timezone.utc).isoformat()
    chunk_ids = [f"{new_doc_id}_{i}" for i in range(len(chunks))]
    metadatas = [
        {
            "doc_id": new_doc_id,
            "filename": file.filename,
            "coach_id": str(current_user.id),
            "client_id": old_doc["client_id"],
            "uploaded_at": uploaded_at,
            "chunk_index": i,
        }
        for i in range(len(chunks))
    ]

    # Delete old Chroma chunks, ingest new ones
    chroma_delete_coach_doc(tenant_id=tenant_id, doc_id=doc_id)
    try:
        chroma_ingest_coach_docs(
            tenant_id=tenant_id, documents=chunks, ids=chunk_ids, metadatas=metadatas
        )
    except Exception:
        logger.exception(
            "Coach document replacement ingestion failed tenant_id=%s coach_id=%s doc_id=%s",
            tenant_id,
            current_user.id,
            doc_id,
        )
        raise HTTPException(status_code=500, detail="Document replacement failed during indexing.")

    # Archive old DB record, persist new one via MCP tools
    await pg_archive_coach_document(doc_id=doc_id)
    await pg_save_coach_document(
        doc_id=new_doc_id,
        coach_id=str(current_user.id),
        client_id=old_doc["client_id"],
        tenant_id=str(current_user.tenant_id),
        filename=file.filename,
        content_type=content_type,
        file_bytes_b64=base64.b64encode(file_bytes).decode(),
        uploaded_at=uploaded_at,
    )

    logger.info(
        "Coach %s replaced doc_id=%s with new_doc_id=%s filename=%s client=%s",
        current_user.id,
        doc_id,
        new_doc_id,
        file.filename,
        old_doc["client_id"],
    )
    return IngestResponse(
        ingested=len(chunks), collection=f"tenant_{tenant_id}_coach_conversations"
    )


@router.delete("/coach/documents/{doc_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_coach_uploaded_document(
    doc_id: str,
    current_user: User = Depends(get_current_user),
    tenant_id: str = Depends(get_tenant_id),
) -> None:
    """Delete a coach-conversation document. Archives DB record and removes Chroma chunks. Coach only."""
    _require_coach(current_user)
    await _get_active_coach_doc(current_user, doc_id)  # ownership check

    chroma_delete_coach_doc(tenant_id=tenant_id, doc_id=doc_id)
    await pg_archive_coach_document(doc_id=doc_id)

    logger.info(
        "Coach %s deleted doc_id=%s tenant=%s", current_user.id, doc_id, tenant_id
    )
